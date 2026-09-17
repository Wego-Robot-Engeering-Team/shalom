#!/usr/bin/env python3
# Copyright (c) 2026 WeGo Robotics. All rights reserved.

"""hmi/resources/error_codes.json 에서 알람 코드 표 엑셀 파일을 만든다.

납품 문서「비상대응 매뉴얼」의 근거 자료이며, 현장에서 코드를 찾아보는 용도다.
카탈로그가 단일 출처이므로 이 파일을 손으로 고치지 말 것.

    python3 tools/gen_alarm_xlsx.py [출력경로]
"""

import collections
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
CATALOG = ROOT / "hmi/resources/error_codes.json"
DEFAULT_OUT = ROOT / "docs/alarm_reference.xlsx"

BANDS = [
    ("E_", "명령 거절"), ("LINK_", "전송"), ("ROBOT_", None), ("SYS_", "시스템"),
    ("TF_", "좌표계"), ("MAP_", "지도"), ("UPDATE_", "배포"), ("SAFETY_", "안전"),
    ("MOTION_", "모션 권한"), ("NAV_", "자율주행"), ("MISSION_", "점검"),
    ("SETUP_", "사전 교시"), ("PERCEPT_", "인식"), ("ARM_", "로봇팔"),
    ("BASE_", "본체"), ("CAPTURE_", "촬영"), ("POWER_", "전원"),
    ("AUDIT_", "조작 이력"),
]
SEVERITY = {"critical": "치명", "error": "오류", "warn": "경고",
            "info": "정보", "ok": "정상"}
ORDER = {"critical": 0, "error": 1, "warn": 2, "info": 3, "ok": 4}

# 등급별 행 색. 조작자가 표를 훑을 때 치명과 오류가 먼저 눈에 들어와야 한다.
ROW_FILL = {
    "치명": PatternFill("solid", fgColor="F4CCCC"),
    "오류": PatternFill("solid", fgColor="FCE5CD"),
    "경고": PatternFill("solid", fgColor="FFF2CC"),
    "정보": PatternFill("solid", fgColor="EFEFEF"),
    "정상": PatternFill("solid", fgColor="D9EAD3"),
}
HEAD_FILL = PatternFill("solid", fgColor="333333")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=10)
MONO = Font(name="Consolas", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("번호", 8), ("코드", 30), ("등급", 10), ("해제", 8), ("영역", 12),
    ("발생처", 12), ("메시지", 34), ("원인", 58), ("조치", 62), ("채널", 22),
    ("억제(초)", 9),
]


def band_of(code):
    return max((p for p, _ in BANDS if code.startswith(p)), key=len)


def rows_from(catalog):
    """카탈로그를 표 순서대로 편다. 영역 순서 → 등급 순서 → 번호 순서."""
    buckets = collections.defaultdict(list)
    for c in catalog["codes"]:
        buckets[band_of(c["code"])].append(c)

    out = []
    for prefix, name in BANDS:
        if name is None:
            continue
        group = list(buckets.get(prefix, []))
        if prefix == "LINK_":           # ROBOT_SELECTED 는 전송에 둔다
            group += buckets.get("ROBOT_", [])
        for c in sorted(group, key=lambda x: (ORDER[x["severity"]], x["id"])):
            clears = {"latched": "유지", "clears_with": "짝 해제"}.get(
                c.get("clears"), "1회")
            out.append({
                "번호": c["id"],
                "코드": c["code"],
                "등급": SEVERITY[c["severity"]],
                "해제": clears,
                "영역": name,
                "발생처": c["origin"],
                "메시지": c["title"],
                "원인": c["cause"],
                "조치": "\n".join(c["actions"]) if c["actions"] else "—",
                "채널": "" if c.get("channel") in (None, "-") else c["channel"],
                "억제(초)": c.get("throttle_s", ""),
            })
    return out


def write_codes(ws, rows):
    ws.freeze_panes = "A2"
    for i, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22

    for r, row in enumerate(rows, start=2):
        for i, (title, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=i, value=row[title])
            cell.font = MONO if title in ("코드", "채널") else BODY
            cell.fill = ROW_FILL[row["등급"]]
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=title in ("메시지", "원인", "조치"),
                horizontal="center" if title in ("번호", "등급", "해제", "억제(초)")
                else "left")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def write_guide(ws, rows, catalog):
    counts = collections.Counter(r["등급"] for r in rows)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 10

    lines = [
        ("제목", "GTX-A 철도차량 하부점검 로봇 — 알람 코드 표", ""),
        ("", "", ""),
        ("출처", "hmi/resources/error_codes.json (규격 버전 "
                 f"{catalog.get('version', '?')})", ""),
        ("생성", "python3 tools/gen_alarm_xlsx.py", ""),
        ("주의", "이 파일을 직접 고치지 말 것. 카탈로그를 고치고 다시 생성한다 — "
                 "손으로 고치면 화면이 띄우는 내용과 어긋난다.", ""),
        ("", "", ""),
        ("등급", "뜻", "개수"),
        ("치명", "운용을 멈추고 조치해야 한다", counts.get("치명", 0)),
        ("오류", "해당 동작이 실패했다. 원인을 없애야 이어갈 수 있다", counts.get("오류", 0)),
        ("경고", "동작은 이어지나 확인이 필요하다", counts.get("경고", 0)),
        ("정보", "알림. 조치가 필요 없다", counts.get("정보", 0)),
        ("정상", "정상 동작 알림", counts.get("정상", 0)),
        ("", "합계", len(rows)),
        ("", "", ""),
        ("해제", "뜻", ""),
        ("유지", "원인이 사라져도 화면에 남는다. 사람이 확인하고 해제해야 한다", ""),
        ("짝 해제", "짝이 되는 코드가 오면 사라진다", ""),
        ("1회", "한 번 알리고 끝난다", ""),
        ("", "", ""),
        ("번호 대역", "영역", ""),
        ("1000", "명령 거절 · 전송 · 시스템 · 좌표계 · 지도 · 배포", ""),
        ("2000", "안전 · 모션 권한", ""),
        ("3000", "자율주행 · 점검 · 사전 교시", ""),
        ("4000", "인식", ""),
        ("5000", "로봇팔 · 본체", ""),
        ("6000", "촬영", ""),
        ("7000", "전원", ""),
        ("8000", "조작 이력", ""),
        ("", "", ""),
        ("억제(초)", "같은 코드가 이 간격 안에 여러 번 나면 한 번만 올린다. "
                     "비어 있으면 올 때마다 올린다.", ""),
    ]
    for r, (a, b, c) in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=a).font = Font(name="Arial", size=10, bold=bool(a))
        ws.cell(row=r, column=2, value=b).font = BODY
        if c != "":
            cell = ws.cell(row=r, column=3, value=c)
            cell.font, cell.alignment = BODY, Alignment(horizontal="center")
    ws["B1"].font = Font(name="Arial", size=14, bold=True)
    for r in range(8, 13):
        ws.cell(row=r, column=1).fill = ROW_FILL[ws.cell(row=r, column=1).value]


def main():
    out = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    rows = rows_from(catalog)

    wb = Workbook()
    write_guide(wb.active, rows, catalog)
    wb.active.title = "읽는 법"
    write_codes(wb.create_sheet("알람 코드"), rows)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{out} 생성 ({len(rows)} 개)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
